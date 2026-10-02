"""Author original synthetic documents with established document libraries."""

import io
import xml.etree.ElementTree as ET
import zipfile
from datetime import datetime
from pathlib import Path


def stable_archive(path: Path) -> None:
    with zipfile.ZipFile(path) as source:
        members = [(name, source.read(name)) for name in sorted(source.namelist())]
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in members:
            if name == "docProps/core.xml":
                core = ET.fromstring(data)
                modified = core.find("{http://purl.org/dc/terms/}modified")
                if modified is not None:
                    modified.text = "2026-10-02T00:00:00Z"
                data = ET.tostring(core, encoding="utf-8")
            info = zipfile.ZipInfo(name, date_time=(2026, 10, 2, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            archive.writestr(info, data)
    path.write_bytes(output.getvalue())


def drawing(path: Path, discipline: str, revision: str) -> None:
    from reportlab.pdfgen.canvas import Canvas

    canvas = Canvas(str(path), pagesize=(600, 420), invariant=1, pageCompression=0)
    canvas.setTitle(f"Golden Project: {discipline} {revision}")
    canvas.setAuthor("Concord synthetic fixture")
    canvas.setFont("Helvetica", 10)
    canvas.drawString(30, 385, "SYNTHETIC COORDINATION FIXTURE / NOT FOR CONSTRUCTION")
    canvas.drawString(30, 365, f"{discipline.upper()} / LEVEL 01 / {revision}")
    canvas.setLineWidth(1)
    canvas.rect(70, 150, 400, 60 if revision == "R1" else 120)
    canvas.setFont("Helvetica", 12)
    canvas.drawString(
        75,
        300,
        "BEAM-01: depth 300 mm"
        if revision == "R1"
        else "BEAM-01: depth 900 mm / Design Change 023",
    )
    canvas.setDash(4, 3)
    route_y = 115 if revision != "R3" else 80
    canvas.line(130, route_y, 350, route_y)
    canvas.setDash()
    canvas.drawString(130, route_y - 20, "DUCT-01 / 400 mm")
    canvas.rect(30, 25, 540, 40)
    canvas.drawString(40, 40, "Drawing legend: solid = beam; dashed = duct route")
    canvas.save()


def specification(path: Path) -> None:
    from reportlab.pdfgen.canvas import Canvas

    canvas = Canvas(str(path), pagesize=(600, 420), invariant=1, pageCompression=0)
    canvas.setTitle("Golden coordination specification")
    canvas.setAuthor("Concord synthetic fixture")
    for page, lines in enumerate(
        [
            [
                "Synthetic coordination specification",
                "Level 01 beam BEAM-01 and duct DUCT-01",
                "R1 beam depth: 300 mm. Maintain duct clearance.",
                "Source marker: GOLDEN-SPEC-PAGE-1",
            ],
            [
                "Review rules",
                "R2 design-change 023 deepens the beam to 900 mm.",
                "Verify rerouted DUCT-01 against the revised structure.",
                "Source marker: GOLDEN-SPEC-PAGE-2",
            ],
        ],
        1,
    ):
        canvas.setFont("Helvetica", 12)
        for index, line in enumerate(lines):
            canvas.drawString(35, 355 - index * 35, line)
        canvas.drawString(35, 40, f"Synthetic original / page {page}")
        canvas.showPage()
    canvas.save()


def workbook(path: Path) -> None:
    from openpyxl import Workbook

    book = Workbook()
    book.properties.creator = "Concord synthetic fixture"
    book.properties.created = datetime(2026, 10, 2)
    book.properties.modified = datetime(2026, 10, 2)
    sheet = book.active
    sheet.title = "Coordination"
    sheet.append(["Element", "Design change", "Status", "Revision"])
    sheet.append(["BEAM-01", "023", "Beam enlarged", "R2"])
    sheet.append(["DUCT-01", "023", "Route requires review", "R2"])
    sheet.append(["DUCT-01", "023", "Route moved; verify clearance", "R3"])
    checks = book.create_sheet("Requirements")
    checks.append(["Element", "Required discipline"])
    checks.append(["BEAM-01", "structure"])
    checks.append(["DUCT-01", "mep"])
    book.save(path)
    stable_archive(path)


def change_document(path: Path) -> None:
    from docx import Document

    document = Document()
    document.core_properties.author = "Concord synthetic fixture"
    document.core_properties.created = datetime(2026, 10, 2)
    document.core_properties.modified = datetime(2026, 10, 2)
    document.add_heading("Design Change 023", level=1)
    document.add_paragraph("Increase BEAM-01 depth from 300 mm to 900 mm at Level 01.")
    document.add_paragraph(
        "DUCT-01 requires a coordinated reroute and targeted clash verification."
    )
    table = document.add_table(rows=1, cols=3)
    for cell, text in zip(table.rows[0].cells, ["Element", "Revision", "Instruction"], strict=True):
        cell.text = text
    for values in [("BEAM-01", "R2", "Enlarge beam"), ("DUCT-01", "R3", "Reroute then verify")]:
        for cell, text in zip(table.add_row().cells, values, strict=True):
            cell.text = text
    document.save(path)
    stable_archive(path)


def scanned_notice(path: Path, font_path: Path) -> None:
    from PIL import Image, ImageDraw, ImageFont

    if not font_path.is_file():
        raise FileNotFoundError("Provide a local Chinese-capable font for fixture authoring")
    image = Image.new("RGB", (1200, 460), "white")
    ink = ImageDraw.Draw(image)
    font = ImageFont.truetype(str(font_path), 42)
    for index, line in enumerate(
        [
            "设计变更 023",
            "梁高度增加，机电管线需要调整",
            "BEAM-01 / DUCT-01 / R2",
            "合成测试数据，不作为施工指令",
        ]
    ):
        ink.text((45, 35 + index * 95), line, fill="black", font=font)
    image.save(path)


def dxf(path: Path, revision: str) -> None:
    import ezdxf

    previous = ezdxf.options.write_fixed_meta_data_for_testing
    ezdxf.options.write_fixed_meta_data_for_testing = True
    document = ezdxf.new("R2010")
    document.layers.new("STRUCTURE")
    document.layers.new("MEP")
    space = document.modelspace()
    space.add_line((0, 0), (4, 0.3 if revision == "R1" else 0.9), dxfattribs={"layer": "STRUCTURE"})
    space.add_line((1, 0.1), (3, 0.1), dxfattribs={"layer": "MEP"})
    document.classes.add_required_classes(document.dxfversion)
    document.classes.classes = dict(sorted(document.classes.classes.items()))
    try:
        document.saveas(path)
    finally:
        ezdxf.options.write_fixed_meta_data_for_testing = previous


def requirements_ids(path: Path) -> None:
    """Write a deterministic IDS file for the Golden Engineering fixture."""
    from ifctester import ids
    from ifctester.ids import Attribute, Entity, Specification

    beam = Specification(name="Beam naming", minOccurs=0, ifcVersion=["IFC4"])
    beam.applicability = [Entity(name="IfcBeam")]
    beam.requirements = [Attribute(name="Name", value="BEAM-01")]
    duct = Specification(name="Duct naming", minOccurs=0, ifcVersion=["IFC4"])
    duct.applicability = [Entity(name="IfcDuctSegment")]
    duct.requirements = [Attribute(name="Name", value="DUCT-01")]
    document = ids.Ids(title="Golden coordination requirements")
    document.specifications.extend((beam, duct))
    path.write_text(document.to_string(), encoding="utf-8")
